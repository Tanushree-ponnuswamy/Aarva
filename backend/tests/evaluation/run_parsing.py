"""Tests E (document reading).  Usage: python run_parsing.py <textbook.pdf> <work-dir> <out.json>"""
import os, sys, json, io, re, time
pdf_path, work, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(work, exist_ok=True)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
import fitz
from services.document_parser import parse_document

res = {}

# 1. Extraction completeness on the real textbook: our parser vs. an independent word count straight from the PDF
data = open(pdf_path, "rb").read()
pr = parse_document(data, "book.pdf")
doc = fitz.open(stream=data, filetype="pdf")
raw_words = sum(len(p.get_text("words")) for p in doc)
try:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pl:
        plumber_words = sum(len((pg.extract_text() or "").split()) for pg in pl.pages)
except Exception:
    plumber_words = None
res["extraction_completeness"] = {"pages": pr.page_count, "words_extracted_by_aarva": pr.word_count, "words_counted_by_pymupdf_word_list": raw_words,
                                  "words_counted_by_pdfplumber": plumber_words,
                                  "captured_percent_vs_pymupdf": round(100 * pr.word_count / raw_words, 1) if raw_words else None,
                                  "captured_percent_vs_pdfplumber": round(100 * pr.word_count / plumber_words, 1) if plumber_words else None,
                                  "empty_pages": sum(1 for p in pr.pages if len((p.text or "").split()) < 5)}

# 2. Format coverage: one small file of each type with a known sentence
SENT = "Photosynthesis lets green plants make glucose from sunlight."
files = {}
d = fitz.open(); p = d.new_page(); p.insert_text((72, 72), SENT); files["sample.pdf"] = d.tobytes()
from docx import Document
dx = Document(); dx.add_paragraph(SENT)
t = dx.add_table(rows=3, cols=2)
for r, (a, b) in enumerate([("Planet", "Moons"), ("Earth", "1"), ("Mars", "2")]):
    t.cell(r, 0).text, t.cell(r, 1).text = a, b
buf = io.BytesIO(); dx.save(buf); files["sample.docx"] = buf.getvalue()
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.append(["Planet", "Moons"]); ws.append(["Earth", 1]); ws.append(["Mars", 2]); ws.append([SENT, ""])
buf = io.BytesIO(); wb.save(buf); files["sample.xlsx"] = buf.getvalue()
files["sample.txt"] = SENT.encode()
from PIL import Image, ImageDraw, ImageFont
img = Image.new("RGB", (1400, 120), "white"); dr = ImageDraw.Draw(img)
try:
    font = ImageFont.truetype("arial.ttf", 40)
except Exception:
    font = ImageFont.load_default()
dr.text((20, 30), SENT, fill="black", font=font)
buf = io.BytesIO(); img.save(buf, "PNG"); files["sample.png"] = buf.getvalue()

ocr_engines = {}
for mod in ("easyocr", "pytesseract"):
    try:
        __import__(mod); ocr_engines[mod] = "installed"
    except Exception:
        ocr_engines[mod] = "not installed"
res["ocr_engines_on_this_machine"] = ocr_engines

fmt = {}
for name, blob in files.items():
    try:
        r = parse_document(blob, name)
        text = (r.full_text or "").lower()
        fmt[name] = {"read_ok": "photosynthesis" in text and "glucose" in text, "words": r.word_count, "warnings": r.parse_warnings[:2]}
    except Exception as e:
        fmt[name] = {"read_ok": False, "error": f"{type(e).__name__}: {e}"}
res["format_coverage"] = fmt

# 3. Tables: are rows readable (planet next to its number)?
tables = {}
for name in ("sample.docx", "sample.xlsx"):
    text = re.sub(r"\s+", " ", parse_document(files[name], name).full_text or "")
    tables[name] = {"earth_row_readable": bool(re.search(r"Earth\W{0,5}1\b", text)), "mars_row_readable": bool(re.search(r"Mars\W{0,5}2\b", text)), "excerpt": text[:160]}
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table
    buf = io.BytesIO(); SimpleDocTemplate(buf, pagesize=A4).build([Table([["Planet", "Moons"], ["Earth", "1"], ["Mars", "2"]])])
    text = re.sub(r"\s+", " ", parse_document(buf.getvalue(), "table.pdf").full_text or "")
    tables["table.pdf"] = {"earth_row_readable": bool(re.search(r"Earth\W{0,5}1\b", text)), "mars_row_readable": bool(re.search(r"Mars\W{0,5}2\b", text)), "excerpt": text[:160]}
except Exception as e:
    tables["table.pdf"] = {"error": f"{type(e).__name__}: {e}"}
res["tables"] = tables

# 4. Scanned (image-only) PDF: 12 pages that contain only a picture of text
d = fitz.open()
for i in range(12):
    p = d.new_page(); p.insert_image(fitz.Rect(40, 40, 560, 140), stream=files["sample.png"])
t0 = time.time(); r = parse_document(d.tobytes(), "scan.pdf")
res["scanned_pdf_12_pages"] = {"words_extracted": r.word_count, "seconds": round(time.time() - t0, 1),
                               "note": "OCR is only attempted for scanned PDFs of 10 pages or fewer", "warnings": r.parse_warnings[:3]}
json.dump(res, open(out_path, "w"), indent=1)
print("PARSING_DONE", json.dumps(res)[:3000], flush=True)
